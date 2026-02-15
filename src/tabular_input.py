from __future__ import annotations

import csv
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator, TextIO


def _stringify_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    return str(value)


def _trim_trailing_empty(values: list[str]) -> list[str]:
    idx = len(values)
    while idx > 0 and not str(values[idx - 1]).strip():
        idx -= 1
    return values[:idx]


def _is_effectively_empty_row(values: list[str]) -> bool:
    return not any(str(v).strip() for v in values)


@contextmanager
def open_input_rows(
    input_path: Path,
) -> Iterator[tuple[list[str], Iterator[tuple[int, list[str]]]]]:
    """
    Open an input file (CSV or XLSX) and yield headers + an iterator of (row_number, row_values).

    Row numbers start at 2 to match the "first data row after header" convention used throughout the app.
    """
    suffix = input_path.suffix.lower()
    if suffix == ".csv":
        f: TextIO | None = None
        try:
            f = input_path.open("r", encoding="utf-8", newline="")
            reader = csv.reader(f)
            headers = next(reader, None)
            if headers is None:
                raise ValueError("Input CSV missing header row.")

            def _iter() -> Iterator[tuple[int, list[str]]]:
                for row_number, row in enumerate(reader, start=2):
                    yield row_number, row

            yield headers, _iter()
        finally:
            if f is not None:
                f.close()
        return

    if suffix == ".xlsx":
        # openpyxl is a small-ish dependency and keeps us off of pandas.
        from openpyxl import load_workbook  # type: ignore[import-not-found]

        wb = None
        try:
            wb = load_workbook(input_path, read_only=True, data_only=True)
            ws = wb.active
            row_iter = ws.iter_rows(values_only=True)
            header_values = next(row_iter, None)
            if header_values is None:
                raise ValueError("Input XLSX missing header row.")

            headers_raw = [_stringify_cell(v).strip() for v in header_values]
            headers_raw = _trim_trailing_empty(headers_raw)
            headers: list[str] = []
            for i, h in enumerate(headers_raw, start=1):
                headers.append(h if h else f"Column {i}")
            if not headers:
                raise ValueError("Input XLSX missing header values.")

            header_len = len(headers)

            def _iter() -> Iterator[tuple[int, list[str]]]:
                for excel_row, values in enumerate(row_iter, start=2):
                    row = [_stringify_cell(v).strip() for v in values]
                    row = _trim_trailing_empty(row)
                    if _is_effectively_empty_row(row):
                        continue
                    if len(row) < header_len:
                        row = row + [""] * (header_len - len(row))
                    elif len(row) > header_len:
                        extras = row[header_len:]
                        if any(v.strip() for v in extras):
                            # Keep parity with CSV behavior (mismatched column counts).
                            yield excel_row, row
                            continue
                        row = row[:header_len]
                    yield excel_row, row

            yield headers, _iter()
        finally:
            if wb is not None:
                wb.close()
        return

    raise ValueError(f"Unsupported input format: {input_path.suffix or '(no extension)'} (supported: .csv, .xlsx)")


def count_data_rows(input_path: Path) -> int:
    """Count data rows (excluding header) for CSV/XLSX inputs."""
    headers: list[str]
    rows: Iterator[tuple[int, list[str]]]
    with open_input_rows(input_path) as (headers, rows):
        # For XLSX we already skip effectively-empty rows; for CSV we count everything after the header.
        _ = headers
        return sum(1 for _ in rows)


def preview_rows(input_path: Path, *, limit: int = 30) -> dict[str, object]:
    if not input_path.exists():
        return {"headers": [], "rows": []}
    headers: list[str]
    rows_iter: Iterator[tuple[int, list[str]]]
    rows: list[list[str]] = []
    with open_input_rows(input_path) as (headers, rows_iter):
        for _, row in zip(range(limit), (r for _, r in rows_iter)):
            rows.append(row)
    return {"headers": headers, "rows": rows}

